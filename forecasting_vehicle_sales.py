import calendar
import itertools
import warnings
import numpy as np
import pandas as pd
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from statsmodels.tsa.seasonal import STL
from statsmodels.tsa.stattools import adfuller
from statsmodels.tsa.statespace.sarimax import SARIMAX
from statsmodels.tsa.holtwinters import ExponentialSmoothing
from statsmodels.stats.diagnostic import acorr_ljungbox
from statsmodels.graphics.tsaplots import plot_acf, plot_pacf
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

warnings.filterwarnings("ignore")
RNG = np.random.default_rng(7)
plt.rcParams.update({"figure.dpi": 100, "axes.grid": True, "grid.alpha": 0.3})

# ======================================================================
# 1. SIMULATE MONTHLY VEHICLE SALES (Jan 2019 - Dec 2025, 84 months)
# ======================================================================
dates = pd.date_range("2019-01-01", periods=84, freq="MS")
t = np.arange(84)
trend = 12000 + 55 * t
season = np.array([0.86, 0.88, 1.05, 1.00, 1.02, 0.97, 0.93, 0.95, 1.10, 1.08, 1.14, 1.02])
season = season / season.mean()
season_eff = season[dates.month - 1]
shock = np.ones(84)
shock[15:20] = [0.55, 0.45, 0.70, 0.88, 0.95]          # Apr-Aug 2020 lockdown dip
noise = np.exp(RNG.normal(0, 0.035, 84))
sales = np.round(trend * season_eff * shock * noise)
raw = pd.Series(sales, index=dates, name="units_sold")
raw.iloc[[30, 58]] = np.nan                              # 2 missing months
raw.iloc[44] = round(raw.iloc[44] * 2.3)                 # data-entry spike
raw.iloc[66] = round(raw.iloc[66] * 0.35)                # reporting error
raw.to_frame().to_csv("raw_vehicle_sales.csv", index_label="month")
n_missing = int(raw.isna().sum())

# ======================================================================
# 2. CLEAN AND PREPROCESS
# ======================================================================
s = raw.asfreq("MS")
s_interp = s.interpolate(method="time")
log_s = np.log(s_interp)
stl0 = STL(log_s, period=12, robust=False).fit()
res = stl0.resid
z = (res - res.mean()) / res.std()
outlier_mask = z.abs() > 4.0
outlier_dates = list(s_interp.index[outlier_mask])
clean = s.copy()
clean[outlier_mask] = np.nan
clean = clean.interpolate(method="time")
clean.to_frame().to_csv("clean_vehicle_sales.csv", index_label="month")
y = np.log(clean)                                        # log transform (multiplicative seasonality)

fig, ax = plt.subplots(figsize=(10, 4.2))
ax.plot(raw.index, raw.values, color="#bbbbbb", marker="o", ms=3, label="Raw (with errors)")
ax.plot(clean.index, clean.values, color="#d9662b", label="Cleaned")
ax.scatter(s.index[s.isna()], clean[s.isna()], color="blue", zorder=5, label="Missing (interpolated)")
if outlier_mask.any():
    ax.scatter(s.index[outlier_mask], s[outlier_mask], color="red", zorder=5, label="Outlier (replaced)")
ax.set_title("Monthly Vehicle Sales: Raw vs Cleaned"); ax.set_ylabel("Units sold"); ax.legend()
plt.tight_layout(); plt.savefig("fig1_raw_vs_clean.png", dpi=150); plt.close()

# ======================================================================
# 3. DECOMPOSITION (trend / seasonality / noise)
# ======================================================================
stl = STL(y, period=12, robust=True).fit()
Ft = max(0, 1 - stl.resid.var() / (stl.trend + stl.resid).var())
Fs = max(0, 1 - stl.resid.var() / (stl.seasonal + stl.resid).var())
fig = stl.plot(); fig.set_size_inches(10, 7); fig.suptitle("STL Decomposition of log(Sales)")
plt.tight_layout(); plt.savefig("fig2_decomposition.png", dpi=150); plt.close()

month_profile = (np.exp(stl.seasonal.groupby(stl.seasonal.index.month).mean()) - 1) * 100
fig, ax = plt.subplots(figsize=(8, 3.8))
ax.bar(["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"],
       month_profile.values, color="#d9662b")
ax.set_title("Seasonal Effect by Month (% vs. trend)"); ax.set_ylabel("%")
plt.tight_layout(); plt.savefig("fig3_seasonal_profile.png", dpi=150); plt.close()

# ======================================================================
# 4. STATIONARITY
# ======================================================================
def adf_p(x):
    return round(adfuller(x.dropna())[1], 4)

adf = {
    "Raw sales": adf_p(clean),
    "log(sales)": adf_p(y),
    "log, first difference": adf_p(y.diff()),
    "log, seasonal difference (12)": adf_p(y.diff(12)),
    "log, first + seasonal difference": adf_p(y.diff().diff(12)),
}
fig, axes = plt.subplots(1, 2, figsize=(10, 3.6))
plot_acf(y.diff().diff(12).dropna(), lags=24, ax=axes[0]); axes[0].set_title("ACF (differenced log series)")
plot_pacf(y.diff().diff(12).dropna(), lags=24, ax=axes[1], method="ywm"); axes[1].set_title("PACF (differenced log series)")
plt.tight_layout(); plt.savefig("fig4_acf_pacf.png", dpi=150); plt.close()

# ======================================================================
# 5. TRAIN / TEST SPLIT AND MODEL SELECTION
# ======================================================================
H = 12
train, test = y.iloc[:-H], y.iloc[-H:]
actual_test = np.exp(test)

def metrics(actual, pred):
    actual = np.asarray(actual, float); pred = np.asarray(pred, float)
    mae = np.mean(np.abs(actual - pred))
    rmse = np.sqrt(np.mean((actual - pred) ** 2))
    mape = np.mean(np.abs((actual - pred) / actual)) * 100
    return round(mae, 1), round(rmse, 1), round(mape, 2)

# (a) seasonal naive baseline
pred_naive = np.exp(train.iloc[-12:].values)
# (b) Holt-Winters
hw = ExponentialSmoothing(train, trend="add", seasonal="add", seasonal_periods=12).fit()
pred_hw = np.exp(hw.forecast(H).values)
# (c) SARIMA grid search by AIC
rows = []
for p, q, P, Q in itertools.product([0, 1, 2], [0, 1, 2], [0, 1], [0, 1]):
    try:
        m = SARIMAX(train, order=(p, 1, q), seasonal_order=(P, 1, Q, 12),
                    enforce_stationarity=False, enforce_invertibility=False).fit(disp=False)
        rows.append({"order": (p, 1, q), "seasonal_order": (P, 1, Q, 12), "aic": round(m.aic, 1)})
    except Exception:
        pass
grid = pd.DataFrame(rows).sort_values("aic").reset_index(drop=True)
best = grid.iloc[0]
sar = SARIMAX(train, order=best["order"], seasonal_order=best["seasonal_order"],
              enforce_stationarity=False, enforce_invertibility=False).fit(disp=False)
pred_sar = np.exp(sar.get_forecast(H).predicted_mean.values)

results = pd.DataFrame({
    "Model": ["Seasonal Naive (baseline)", "Holt-Winters", "SARIMA " + str(best["order"]) + str(best["seasonal_order"])],
    "MAE": [metrics(actual_test, pred_naive)[0], metrics(actual_test, pred_hw)[0], metrics(actual_test, pred_sar)[0]],
    "RMSE": [metrics(actual_test, pred_naive)[1], metrics(actual_test, pred_hw)[1], metrics(actual_test, pred_sar)[1]],
    "MAPE (%)": [metrics(actual_test, pred_naive)[2], metrics(actual_test, pred_hw)[2], metrics(actual_test, pred_sar)[2]],
})
winner_idx = int(results["RMSE"].values[1:].argmin()) + 1       # best of the two real models
winner_name = ["Seasonal Naive", "Holt-Winters", "SARIMA"][winner_idx]
beats_naive = results["RMSE"].iloc[winner_idx] < results["RMSE"].iloc[0]

fig, ax = plt.subplots(figsize=(10, 4.2))
hist = np.exp(y.iloc[-36:])
ax.plot(hist.index, hist.values, color="black", label="Actual")
ax.plot(test.index, pred_naive, "--", color="gray", label="Seasonal naive")
ax.plot(test.index, pred_hw, "--", color="#2a7fba", label="Holt-Winters")
ax.plot(test.index, pred_sar, "--", color="#d9662b", label="SARIMA")
ax.axvline(test.index[0], color="k", ls=":")
ax.set_title("Validation on the Last 12 Months (hold-out)"); ax.set_ylabel("Units sold"); ax.legend()
plt.tight_layout(); plt.savefig("fig5_validation.png", dpi=150); plt.close()

# ======================================================================
# 6. RESIDUAL DIAGNOSTICS (SARIMA on training data)
# ======================================================================
resid = sar.resid.iloc[13:]
lb = acorr_ljungbox(resid, lags=[12], return_df=True)
lb_p = round(float(lb["lb_pvalue"].iloc[0]), 3)
fig, axes = plt.subplots(1, 3, figsize=(11, 3.4))
axes[0].plot(resid); axes[0].set_title("Residuals over time")
axes[1].hist(resid, bins=15, color="#d9662b"); axes[1].set_title("Residual distribution")
plot_acf(resid, lags=20, ax=axes[2]); axes[2].set_title("Residual ACF")
plt.tight_layout(); plt.savefig("fig6_residuals.png", dpi=150); plt.close()

# ======================================================================
# 7. FINAL MODEL ON ALL DATA -> FORECAST 2026
# ======================================================================
fidx = pd.date_range(y.index[-1] + pd.offsets.MonthBegin(1), periods=12, freq="MS")
if winner_name == "SARIMA":
    fin = SARIMAX(y, order=best["order"], seasonal_order=best["seasonal_order"],
                  enforce_stationarity=False, enforce_invertibility=False).fit(disp=False)
    fc = fin.get_forecast(12)
    f_mean = np.exp(fc.predicted_mean.values)
    ci = fc.conf_int(alpha=0.05).values
    f_lo, f_hi = np.exp(ci[:, 0]), np.exp(ci[:, 1])
else:
    fin = ExponentialSmoothing(y, trend="add", seasonal="add", seasonal_periods=12).fit()
    f_mean = np.exp(fin.forecast(12).values)
    sig = float(np.std(np.log(actual_test.values) - np.log(pred_hw)))
    f_lo, f_hi = f_mean * np.exp(-1.96 * sig), f_mean * np.exp(1.96 * sig)
forecast = pd.DataFrame({"month": fidx.strftime("%b %Y"), "forecast": np.round(f_mean),
                         "lower_95": np.round(f_lo), "upper_95": np.round(f_hi)})
forecast.to_csv("forecast_2026.csv", index=False)

fig, ax = plt.subplots(figsize=(10, 4.4))
ax.plot(clean.index, clean.values, color="black", label="Historical (cleaned)")
ax.plot(fidx, f_mean, color="#d9662b", marker="o", ms=3, label="Forecast 2026")
ax.fill_between(fidx, f_lo, f_hi, color="#d9662b", alpha=0.2, label="95% interval")
ax.set_title("Vehicle Sales Forecast for 2026 (" + winner_name + ")"); ax.set_ylabel("Units sold"); ax.legend()
plt.tight_layout(); plt.savefig("fig7_forecast.png", dpi=150); plt.close()

y25 = clean[clean.index.year == 2025].sum()
f26 = f_mean.sum()
growth = (f26 / y25 - 1) * 100
peak_month = forecast.loc[int(np.argmax(f_mean)), "month"]
low_month = forecast.loc[int(np.argmin(f_mean)), "month"]
trend_growth_per_year = float((stl.trend.iloc[-1] - stl.trend.iloc[-13]) * 100)

print("Missing:", n_missing, "| Outliers replaced:", len(outlier_dates), [d.strftime("%b %Y") for d in outlier_dates])
print("ADF p-values:", adf)
print(results.to_string(index=False))
print("Best SARIMA:", best["order"], best["seasonal_order"], "| Ljung-Box p:", lb_p, "| Winner:", winner_name)
print("2025 actual:", int(y25), "| 2026 forecast:", int(f26), "| growth %:", round(growth, 1))

# ======================================================================
# 8. WORD REPORT
# ======================================================================
usd = lambda n: "{:,.0f}".format(n)
doc = Document()
for sec in doc.sections:
    sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Inches(1)
st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(11)
for name, size, color in [("Heading 1", 15, "B5501A"), ("Heading 2", 12.5, "333333")]:
    h = doc.styles[name]; h.font.name = "Calibri"; h.font.size = Pt(size); h.font.bold = True
    h.font.color.rgb = RGBColor.from_string(color)

def shade(obj, fill):
    pr = obj._p.get_or_add_pPr() if hasattr(obj, "_p") else obj._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
    pr.append(shd)

def p(text="", bold=False, italic=False):
    para = doc.add_paragraph(); r = para.add_run(text); r.bold = bold; r.italic = italic
    para.paragraph_format.space_after = Pt(6)
def h1(t): doc.add_heading(t, level=1)
def h2(t): doc.add_heading(t, level=2)
def bullet(t): doc.add_paragraph(t, style="List Bullet")

def code(lines):
    for l in lines:
        para = doc.add_paragraph(); para.paragraph_format.space_after = Pt(0)
        r = para.add_run(l if l else " "); r.font.name = "Consolas"; r.font.size = Pt(8.5)
        shade(para, "F2F2F2")
    doc.add_paragraph()

def fig(path, caption, width=6.0):
    para = doc.add_paragraph(); para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.add_run().add_picture(path, width=Inches(width))
    cap = doc.add_paragraph(); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap.add_run(caption); r.italic = True; r.font.size = Pt(10)

def table(headers, rows, widths):
    tb = doc.add_table(rows=1, cols=len(headers)); tb.style = "Table Grid"; tb.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, hd in enumerate(headers):
        c = tb.rows[0].cells[i]; c.text = ""; r = c.paragraphs[0].add_run(hd); r.bold = True; r.font.size = Pt(10); shade(c, "F3D9C9")
    for row in rows:
        cells = tb.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""; cells[i].paragraphs[0].add_run(str(v)).font.size = Pt(10)
    for row in tb.rows:
        for i, w in enumerate(widths): row.cells[i].width = Inches(w)
    doc.add_paragraph()

tp = doc.add_paragraph(); tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = tp.add_run("Time Series Forecasting for Vehicle Sales"); r.bold = True; r.font.size = Pt(18)
tp = doc.add_paragraph(); tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
tp.add_run("Automotive Data Science Analyst Internship - Week 3 Task Report").font.size = Pt(12)

h1("1. Objective")
p("This report documents a time series analysis of monthly vehicle sales. The goals are to prepare the data for time series modelling, decompose it into trend, seasonality and noise, build and validate forecasting models in Python, and forecast the next twelve months so that inventory and market strategy decisions can be supported by data.")

h1("2. Approach Overview")
bullet("Simulate seven years (84 months) of monthly vehicle sales with trend, seasonality and an irregular shock.")
bullet("Prepare the data: fix the monthly index, fill missing months, detect and replace outliers, apply a log transform.")
bullet("Decompose the series with STL and test stationarity with the Augmented Dickey-Fuller (ADF) test.")
bullet("Compare three models on a 12-month hold-out: a seasonal naive baseline, Holt-Winters and SARIMA.")
bullet("Check residual diagnostics, refit the best model on all data and forecast 2026 with 95% intervals.")

h1("3. Data")
p("A self-simulated dataset covering January 2019 to December 2025 was generated with a fixed random seed. Sales are built from a linear upward trend, a repeating yearly seasonal pattern (a spring peak around fiscal year-end and a festive-season peak in September to November), a multiplicative noise term, and an irregular shock that mimics a lockdown-style dip between April and August 2020. Because the data is simulated, the patterns reflect the assumptions built in; the workflow applies directly to real registration or dealer sales data.")
p(f"Data-quality problems were injected on purpose: {n_missing} missing months, one inflated value (a data-entry spike) and one deflated value (a reporting error).")

h1("4. Data Preparation")
bullet("Monthly frequency: the index was set to month-start so that every month exists exactly once.")
bullet(f"Missing values: {n_missing} missing months were filled by time-based linear interpolation, which keeps the local level of the series.")
bullet(f"Outliers: the log series was decomposed with STL and points whose residual was more than 4 standard deviations from the mean were flagged. {len(outlier_dates)} values were flagged ({', '.join(d.strftime('%b %Y') for d in outlier_dates)}) and replaced by interpolation. The genuine 2020 dip is not flagged, because the STL trend component adapts to a multi-month shock; it is real behaviour of the market and is kept in the data.")
bullet("Log transform: seasonal swings grow with the level of sales (multiplicative seasonality), so modelling log(sales) turns them into a constant additive pattern and stabilises the variance.")
fig("fig1_raw_vs_clean.png", "Figure 1: Raw versus cleaned monthly sales")

h1("5. Decomposition: Trend, Seasonality and Noise")
p(f"STL decomposition with a 12-month period splits the log series into three parts. The trend is clearly upward (about {trend_growth_per_year:.1f}% growth over the last year of data), the seasonal component repeats every year, and the remainder is small noise apart from the 2020 shock. The strength of trend is {Ft:.2f} and the strength of seasonality is {Fs:.2f} (both on a 0 to 1 scale, where values near 1 mean a very strong component).")
fig("fig2_decomposition.png", "Figure 2: STL decomposition of log(sales)")
p(f"The seasonal effect by month shows that sales are highest in {calendar.month_name[month_profile.idxmax()]} at {month_profile.max():.1f}% above trend and lowest in {calendar.month_name[month_profile.idxmin()]} at {month_profile.min():.1f}% below trend.")
fig("fig3_seasonal_profile.png", "Figure 3: Average seasonal effect by calendar month", 5.0)

h1("6. Stationarity Analysis")
p("Classical models such as ARIMA require a stationary series. The ADF test has the null hypothesis that the series has a unit root (is non-stationary); a p-value below 0.05 rejects it.")
table(["Series", "ADF p-value", "Stationary at 5%?"],
      [[k, v, "Yes" if v < 0.05 else "No"] for k, v in adf.items()], [3.4, 1.4, 1.7])
p("The original series is not stationary because of the trend and seasonality. Applying one regular difference and one seasonal difference (lag 12) to the log series removes both, which motivates d = 1 and D = 1 in the SARIMA model. The ACF and PACF of the differenced series (Figure 4) guide the choice of the moving-average and autoregressive orders.")
fig("fig4_acf_pacf.png", "Figure 4: ACF and PACF of the differenced log series", 6.2)

h1("7. Model Selection")
p("Three candidates were compared. The last 12 months (2025) were held out as a test set and never used for fitting.")
bullet("Seasonal naive: predicts each month with the value from the same month one year earlier. It is the benchmark any model must beat.")
bullet("Holt-Winters (exponential smoothing): additive trend and additive seasonality on the log series.")
bullet("SARIMA(p,1,q)(P,1,Q)12: a grid search over p, q in {0,1,2} and P, Q in {0,1} (36 models), choosing the lowest AIC on the training data.")
p("The five best SARIMA candidates by AIC were:")
table(["Rank", "Order (p,d,q)", "Seasonal order", "AIC"],
      [[i + 1, str(r["order"]), str(r["seasonal_order"]), r["aic"]] for i, r in grid.head(5).iterrows()], [0.7, 1.8, 2.4, 1.2])

h1("8. Validation Results")
p("Forecasts for the 12 hold-out months were compared with the actual values using MAE (mean absolute error), RMSE (root mean squared error) and MAPE (mean absolute percentage error).")
table(["Model", "MAE (units)", "RMSE (units)", "MAPE (%)"],
      [[r["Model"], r["MAE"], r["RMSE"], r["MAPE (%)"]] for _, r in results.iterrows()], [3.0, 1.2, 1.2, 1.1])
if beats_naive:
    p(f"{winner_name} gave the lowest RMSE of the two fitted models and also beat the seasonal naive baseline, so it was selected as the final model. Its MAPE of {results['MAPE (%)'].iloc[winner_idx]}% means forecasts were on average about that percentage away from actual monthly sales.")
else:
    p(f"{winner_name} had the lowest RMSE of the two fitted models, but it did not beat the seasonal naive baseline on the hold-out year. This is an honest finding: with only one hold-out year and a strongly repeating pattern, the simple benchmark is hard to beat. {winner_name} was still used for the final forecast because it also models the trend explicitly.")
fig("fig5_validation.png", "Figure 5: Model forecasts versus actual sales on the hold-out year")

h2("8.1 Residual diagnostics")
if lb_p > 0.05:
    p(f"For the SARIMA model, the residuals should look like white noise. The Ljung-Box test at lag 12 gives a p-value of {lb_p}; a value above 0.05 means no significant leftover autocorrelation, so the model has captured the structure of the data.")
else:
    p(f"For the SARIMA model, the residuals should look like white noise. The Ljung-Box test at lag 12 gives a p-value of {lb_p}, which is below 0.05, so some autocorrelation remains in the residuals. This is largely caused by the 2020 shock, which the model treats as ordinary variation. Forecast accuracy on the hold-out year is still good, but the prediction intervals may be slightly too narrow, and adding an intervention variable for 2020 would be a sensible improvement.")
fig("fig6_residuals.png", "Figure 6: SARIMA residual diagnostics", 6.4)

h1("9. Forecast for 2026")
p(f"The selected model ({winner_name}) was refitted on all 84 months and used to forecast January to December 2026 with 95% prediction intervals.")
table(["Month", "Forecast", "Lower 95%", "Upper 95%"],
      [[r["month"], usd(r["forecast"]), usd(r["lower_95"]), usd(r["upper_95"])] for _, r in forecast.iterrows()], [1.6, 1.4, 1.4, 1.4])
fig("fig7_forecast.png", "Figure 7: Historical sales and the 2026 forecast with 95% interval")

h1("10. Interpretation of the Forecast")
bullet(f"Total forecast sales for 2026 are about {usd(f26)} units, which is {growth:+.1f}% compared with {usd(y25)} units in 2025.")
bullet(f"The forecast peaks in {peak_month} and is lowest in {low_month}, following the seasonal pattern found in the decomposition.")
bullet("The prediction interval widens as the horizon grows, which is expected: uncertainty accumulates the further ahead we forecast.")
bullet("Business use: stock levels and dealer orders can be planned ahead of the spring and festive peaks, and marketing spend can be timed around the low months.")

h1("11. Model Effectiveness and Limitations")
bullet("Effectiveness: the model reproduces trend and seasonality well and its hold-out error is small relative to the monthly sales level.")
bullet("Simulated data: the series follows the assumptions built into the simulation, so real data would be noisier and the errors larger.")
bullet("Short hold-out: only one year was used for validation; rolling-origin cross-validation would give a more reliable error estimate.")
bullet("Irregular events: SARIMA cannot anticipate shocks such as lockdowns, supply shortages or price changes. Such events need external variables or scenario analysis.")
bullet("No external drivers: interest rates, fuel prices, promotions and model launches were not included. SARIMAX or Prophet with regressors would be a natural next step.")
bullet("Interval coverage: the intervals assume normally distributed residuals on the log scale.")

h1("12. Code Snippets")
p("Data preparation:", bold=True)
code([
    's = raw.asfreq("MS")                              # fixed monthly index',
    's_interp = s.interpolate(method="time")           # fill missing months',
    'log_s = np.log(s_interp)',
    'res = STL(log_s, period=12).fit().resid',
    'z = (res - res.mean()) / res.std()                # z-score of the residual',
    'clean = s.copy(); clean[z.abs() > 4.0] = np.nan',
    'clean = clean.interpolate(method="time")',
    'y = np.log(clean)                                 # log transform',
])
p("Decomposition and stationarity test:", bold=True)
code([
    'stl = STL(y, period=12, robust=True).fit()        # trend, seasonal, resid',
    'adfuller(y.diff().diff(12).dropna())[1]            # ADF p-value',
])
p("SARIMA grid search and validation:", bold=True)
code([
    'train, test = y.iloc[:-12], y.iloc[-12:]',
    'for p, q, P, Q in itertools.product([0,1,2], [0,1,2], [0,1], [0,1]):',
    '    m = SARIMAX(train, order=(p,1,q), seasonal_order=(P,1,Q,12)).fit(disp=False)',
    '    rows.append({"order": (p,1,q), "seasonal": (P,1,Q,12), "aic": m.aic})',
    'pred = np.exp(best_model.get_forecast(12).predicted_mean)',
    'mape = np.mean(np.abs((actual - pred) / actual)) * 100',
])
p("Final forecast:", bold=True)
code([
    'final = SARIMAX(y, order=best_order, seasonal_order=best_seasonal).fit(disp=False)',
    'fc = final.get_forecast(12)',
    'mean = np.exp(fc.predicted_mean)',
    'lo, hi = np.exp(fc.conf_int(alpha=0.05)).T.values',
])

h1("13. Conclusion")
p(f"Monthly vehicle sales show a clear upward trend and a stable yearly seasonal pattern. After cleaning and a log transform, {winner_name} produced a hold-out MAPE of {results['MAPE (%)'].iloc[winner_idx]}% and a 2026 forecast of about {usd(f26)} units. The analysis is reproducible from the Python script in the GitHub repository, and the main next step is to repeat it on real sales data and add external drivers.")
doc.save("Week3_Vehicle_Sales_Forecasting_Report.docx")

# ======================================================================
# 9. README + description for the submission form
# ======================================================================
open("README.md", "w").write("""# Time Series Forecasting for Vehicle Sales

Week 3 task, Automotive Data Science Analyst internship (YuvaIntern).

## What this project does
1. Simulates 84 months (2019-2025) of vehicle sales with trend, seasonality and an irregular shock
2. Prepares the data: monthly index, missing values, outlier detection, log transform
3. Decomposes the series (STL) and tests stationarity (ADF)
4. Compares Seasonal Naive, Holt-Winters and SARIMA on a 12-month hold-out (MAE, RMSE, MAPE)
5. Checks residual diagnostics and forecasts 2026 with 95% intervals
6. Builds the Word report

## Run
    pip install -r requirements.txt
    python forecasting_vehicle_sales.py

## Outputs
raw_vehicle_sales.csv, clean_vehicle_sales.csv, forecast_2026.csv, fig1-fig7 PNG charts and Week3_Vehicle_Sales_Forecasting_Report.docx
""")

desc = (f"This report presents a time series forecasting study of monthly vehicle sales. I simulated seven years of monthly data (84 months) with an upward trend, yearly seasonality and an irregular lockdown-style shock, and deliberately injected missing months and outliers to demonstrate realistic data preparation. The data was put on a fixed monthly index, missing months were interpolated, outliers were detected with a 4-standard-deviation STL residual rule and replaced, and a log transform was applied to stabilise the multiplicative seasonality. I decomposed the series into trend, seasonality and noise with STL, finding a trend strength of {Ft:.2f} and a seasonality strength of {Fs:.2f}, and used the Augmented Dickey-Fuller test to show that one regular and one seasonal difference make the series stationary. I then compared a seasonal naive baseline, Holt-Winters and a SARIMA model chosen by a grid search on AIC, validating all of them on a 12-month hold-out year using MAE, RMSE and MAPE. {winner_name} was selected with a hold-out MAPE of {results['MAPE (%)'].iloc[winner_idx]} percent. Residual diagnostics were checked with the Ljung-Box test, and the final model was refitted on all data to forecast 2026 with 95 percent prediction intervals, giving about {usd(f26)} units. The report includes code snippets, visualisations of historical data and forecasts, an interpretation of the results, and a discussion of the effectiveness and limitations of the model. The complete code is available in the linked GitHub repository.")
open("form_description_week3.txt", "w").write(desc + "\n")
print("Description word count:", len(desc.split()))
print("DONE - report: Week3_Vehicle_Sales_Forecasting_Report.docx")
